"""
ARUNA — deck PPTX ilmiah (13 slide + referensi), video latar looping di setiap slide, Morph + animasi otomatis.

  python design/showcase/deck/deck.py [en|id]
Keluaran: design/showcase/deck/ARUNA_deck_<lang>.pptx
Aset   : assets/bg_loop.mp4 (make_bg_loop.py), assets/hw_*.png + hw_points.json (hw_xray.py),
         assets/cover.png (cover_render.py), ../video/frames (scene.py), ../../app/renders, ../../app/screens
"""
import copy
import glob
import json
import os
import sys

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

LANG = (sys.argv[1] if len(sys.argv) > 1 else os.environ.get("APP_LANG", "en")).lower()
HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, "assets")
SHOW = os.path.dirname(HERE)
DESIGN = os.path.dirname(SHOW)
CACHE = os.path.join(HERE, "_cache")
os.makedirs(CACHE, exist_ok=True)


def T(en, id_):
    return en if LANG == "en" else id_


# ------------------------------------------------------------------ design tokens (sama dengan web/app)
INK, INK2, MUTED = RGBColor(0x10, 0x20, 0x1A), RGBColor(0x4A, 0x5B, 0x52), RGBColor(0x86, 0x95, 0x8D)
GREEN, TEAL, DEEP = RGBColor(0x15, 0x8A, 0x52), RGBColor(0x0E, 0x7A, 0x6A), RGBColor(0x0F, 0x6B, 0x40)
SOFT, AMBER, BLUE, RED = RGBColor(0xE3, 0xF3, 0xEA), RGBColor(0xB8, 0x74, 0x1A), RGBColor(0x2D, 0x6B, 0xD8), RGBColor(0xC8, 0x44, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
FONT_B = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
M = 0.6  # margin (in)


# ------------------------------------------------------------------ kartu PIL (di-bake)
def card_png(w, h, kind="raise", r=0.28):
    key = f"{kind}_{w:.2f}_{h:.2f}_{r:.2f}.png"
    path = os.path.join(CACHE, key)
    if os.path.exists(path):
        return path
    S = 200  # px per inch
    pad = int(0.35 * S)
    W, H, R = int(w * S), int(h * S), int(r * S)
    im = Image.new("RGBA", (W + 2 * pad, H + 2 * pad), (0, 0, 0, 0))
    sh = Image.new("L", im.size, 0)
    ImageDraw.Draw(sh).rounded_rectangle((pad + 10, pad + 22, pad + W + 4, pad + H + 22), R, fill=255)
    sh = sh.filter(ImageFilter.GaussianBlur(26))
    col = (21, 120, 80, 0) if kind != "aura" else (21, 160, 90, 0)
    alpha = 0.16 if kind != "aura" else 0.40
    shadow = Image.new("RGBA", im.size, col[:3] + (0,))
    shadow.putalpha(sh.point(lambda v: int(v * alpha)))
    im.alpha_composite(shadow)
    body = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    grad = Image.new("RGBA", (W, H))
    top, bot = ((255, 255, 255), (244, 250, 246)) if kind != "aura" else ((255, 255, 255), (226, 244, 233))
    for y in range(H):
        t = y / max(1, H - 1)
        c = tuple(int(top[i] * (1 - t) + bot[i] * t) for i in range(3))
        ImageDraw.Draw(grad).line([(0, y), (W, y)], fill=c + (int(255 * 0.95),))
    mask = Image.new("L", (W, H), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, W - 1, H - 1), R, fill=255)
    body.paste(grad, (0, 0), mask)
    ImageDraw.Draw(body).rounded_rectangle((1, 1, W - 2, H - 2), R, outline=(255, 255, 255, 230), width=3)
    im.alpha_composite(body, (pad, pad))
    im.save(path)
    return path


def card(slide, x, y, w, h, kind="raise", r=0.28, name=None):
    pad = 0.35
    pic = slide.shapes.add_picture(card_png(w, h, kind, r), Inches(x - pad), Inches(y - pad),
                                   Inches(w + 2 * pad), Inches(h + 2 * pad))
    if name:
        pic.name = name
    return pic


# ------------------------------------------------------------------ helper teks & bentuk
def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name=None, spacing=1.0, font=FONT):
    """runs: str | [(teks, {opsi})] | [[...baris...]]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, o in line:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
            if o.get("sub"):
                r._r.get_or_add_rPr().set("baseline", "-25000")
    if name:
        tb.name = name
    return tb


def pill(slide, x, y, label, fill=WHITE, color=GREEN, size=11, name=None, w=None, bold=True):
    w = w or (0.22 + len(label) * size * 0.0098)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36))
    s.adjustments[0] = 0.5
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = RGBColor(0xDC, 0xE7, 0xE1) if fill == WHITE else fill
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if name:
        s.name = name
    return s


def dot(slide, x, y, d, fill, label=None, color=WHITE, size=12, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if label:
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
        r.font.color.rgb = color
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if name:
        s.name = name
    return s


def glow(shape, color="158A52", rad=8, alpha=35):
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    g = etree.SubElement(eff, qn("a:glow"), rad=str(int(rad * 12700)))
    c = etree.SubElement(g, qn("a:srgbClr"), val=color)
    etree.SubElement(c, qn("a:alpha"), val=str(alpha * 1000))


def pic(slide, path, x, y, w=None, h=None, name=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if name:
        p.name = name
    return p


def fit_image(src, max_w, max_h, crop=None, key=None):
    """Simpan salinan JPG/PNG yang sudah di-crop & diperkecil (ukuran deck tetap ringan)."""
    im = Image.open(src)
    if crop:
        W, H = im.size
        im = im.crop((int(crop[0] * W), int(crop[1] * H), int(crop[2] * W), int(crop[3] * H)))
    im.thumbnail((int(max_w * 200), int(max_h * 200)), Image.LANCZOS)
    out = os.path.join(CACHE, (key or os.path.basename(src)) + (".png" if im.mode == "RGBA" else ".jpg"))
    if im.mode == "RGBA":
        im.save(out)
    else:
        im.convert("RGB").save(out, quality=88)
    return out, im.size[0] / 200, im.size[1] / 200


# ------------------------------------------------------------------ kerangka slide
SLIDES = []


def new_slide(num, chip, title1, title2, desc=None, title_w=11.6, desc_w=6.4):
    s = prs.slides.add_slide(BLANK)
    v = s.shapes.add_movie(os.path.join(A, "bg_loop.mp4"), 0, 0, prs.slide_width, prs.slide_height,
                           poster_frame_image=os.path.join(A, "bg_poster.png"), mime_type="video/mp4")
    v.name = "BG_VIDEO"
    anim = []
    if chip:
        anim.append(pill(s, M, 0.5, f"{num:02d} · {chip}", name="Chip"))
    if title1:
        anim.append(text(s, M, 0.98, title_w, 1.25, [[(title1 + " ", {"color": INK}), (title2, {"color": GREEN})]],
                         size=32, bold=True, spacing=0.92, name="Title"))
    if desc:
        anim.append(text(s, M, 1.62 if title1 else 1.0, desc_w, 0.6, desc, size=14, color=INK2, name="Desc"))
    # footer
    text(s, M, 7.05, 4, 0.25, [("ARUNA", {"bold": True, "color": GREEN}), ("  ·  smart home biodigester", {"color": MUTED})],
         size=9, name="Footer")
    text(s, 12.2, 7.05, 0.55, 0.25, f"{num:02d}", size=9, color=MUTED, align=PP_ALIGN.RIGHT, name="PageNo")
    SLIDES.append({"slide": s, "video": v, "anim": anim, "float": []})
    return s, SLIDES[-1]


def illustrative(s, meta, x, y):
    meta["anim"].append(pill(s, x, y, T("ILLUSTRATIVE DATA · replace with field results",
                                        "DATA ILUSTRASI · ganti dengan hasil uji"), fill=RGBColor(0xFB, 0xF0, 0xDF),
                             color=AMBER, size=9, name="Tag_Ilustrasi"))


def frame(n):
    p = os.path.join(SHOW, "video", "frames", f"f_{n:04d}.png")
    if not os.path.exists(p):  # render belum selesai: pakai frame terdekat yang ada (build ulang setelah render)
        have = sorted(glob.glob(os.path.join(SHOW, "video", "frames", "f_*.png")))
        print("[DECK] PERINGATAN frame", n, "belum ada")
        p = have[-6]   # hindari frame placeholder/sedang ditulis
    return p


# ================================================================== KONTEN ILMIAH (disisipkan ke deck.py)
# Angka desain: design/competition/calc.json (sci_calc.py). Figur: design/competition/fig/.
COMP = os.path.join(DESIGN, "competition")
CALC = json.load(open(os.path.join(COMP, "calc.json")))
H4, NM, MX = CALC["hh4"], CALC["nominal"], CALC["max"]
SF = os.path.join(COMP, "fig")
NOSTYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
f0 = lambda v: f"{v:,.0f}"
f1 = lambda v: f"{v:,.1f}"
dec = lambda s_: s_.replace(".", ",") if LANG == "id" else s_
FWYR = f"{NM['fw'] * 365 / 1000:.2f}"
CO2N = f"{NM['co2e']:.2f}"


def cite(s, meta, txt):
    meta["anim"].append(text(s, M, 6.74, 11.9, 0.3, [(T("Sources: ", "Sumber: "), {"bold": True}), (txt, {})],
                             size=8.5, color=MUTED, name="Cite"))


def ntable(s, x, y, w, rows, colw, size=10.5, head_fill=DEEP, hl_col=None, name="Tabel", row_h=0.3):
    tb = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tb.name = name
    tblPr = tb._element.graphic.graphicData.tbl.tblPr
    for st in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(st)
    etree.SubElement(tblPr, qn("a:tableStyleId")).text = NOSTYLE
    for j, cw in enumerate(colw):
        tb.table.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.table.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tb.table.cell(i, j)
            c.text = ""
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = FONT, Pt(size if i else size - 0.5)
            r.font.bold = i == 0 or (hl_col is not None and j == hl_col)
            r.font.color.rgb = WHITE if i == 0 else (DEEP if hl_col is not None and j == hl_col else INK)
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if i == 0 else (SOFT if (hl_col is not None and j == hl_col) else
                                                              (WHITE if i % 2 else RGBColor(0xF4, 0xFA, 0xF6)))
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tb


def figure(s, meta, path, x, y, w=None, h=None, name="Fig", pad=0.08, card_kind="raise", flo=False):
    im = Image.open(path)
    ar = im.height / im.width
    if w is None:
        w = h / ar
    h = w * ar
    meta["anim"].append(card(s, x - pad, y - pad, w + 2 * pad, h + 2 * pad, card_kind, r=0.18, name=name + "_card"))
    p = pic(s, path, x, y, w, h, name=("FX_FLOAT_" if flo else "") + name)
    meta["anim"].append(p)
    if flo:
        meta["float"].append(p)
    return h


def bullets(s, meta, x, y, w, h, items, size=12.5, name="Bul", col=GREEN):
    meta["anim"].append(text(s, x, y, w, h, [[("● ", {"color": col, "size": size - 2}), (it[0], {"bold": True}), (it[1], {"color": INK2})]
                                             if isinstance(it, tuple) else [("● ", {"color": col, "size": size - 2}), (it, {"color": INK2})]
                                             for it in items], size=size, spacing=1.08, name=name))


SDGC = {"7": RGBColor(0xC8, 0x9A, 0x08), "11": RGBColor(0xE0, 0x80, 0x10), "12": RGBColor(0xA0, 0x74, 0x26), "13": RGBColor(0x3F, 0x7E, 0x44)}

# ================================================================== 1 COVER
s, m = new_slide(1, None, None, None)
cv, cw, ch = fit_image(os.path.join(A, "cover.png"), 5.6, 6.6, key="cover_sci")
c = pic(s, cv, 13.333 - cw - 0.35, 7.5 - ch - 0.25, cw, ch, name="FX_FLOAT_Cover")
m["float"].append(c)
m["anim"].append(pill(s, M, 0.95, T("RESEARCH & DEVELOPMENT · ENVIRONMENTAL ENGINEERING", "RISET & PENGEMBANGAN · TEKNIK LINGKUNGAN"), name="Chip"))
m["anim"].append(text(s, M, 1.45, 7.6, 1.2, "ARUNA", size=80, bold=True, name="Title"))
m["anim"].append(text(s, M, 2.75, 7.4, 1.5, [[(T("Kinetic-based design of an IoT-enabled ", "Desain berbasis kinetika digester anaerob "), {"color": INK}),
                                              (T("household anaerobic digester for food waste", "rumah tangga ber-IoT untuk sisa makanan"), {"color": GREEN})]],
                      size=24, bold=True, spacing=0.95, name="Sub"))
m["anim"].append(text(s, M, 4.15, 7.0, 0.9, T(
    "Anaerobic Recycling and Upcycling of Nutrient-rich food wAste: sized from mass balance and Gompertz kinetics, instrumented with NDIR methane sensing and safety interlocks.",
    "Anaerobic Recycling and Upcycling of Nutrient-rich food wAste: diukur dari neraca massa dan kinetika Gompertz, dilengkapi sensor metana NDIR dan interlock keamanan."),
    size=13, color=INK2, name="Desc"))
for i, (n, tg) in enumerate([("7", "7.1"), ("11", "11.6"), ("12", "12.3"), ("13", "13.2")]):
    m["anim"].append(pill(s, M + i * 1.12, 5.3, f"SDG {tg}", fill=SDGC[n], color=WHITE, size=10, w=1.0, name=f"SDG_{i}"))
m["anim"].append(text(s, M, 6.0, 7, 0.4, [[("[Team name]", {"bold": True, "color": AMBER}), ("  ·  [Supervisor]  ·  [Institution]  ·  2026", {"color": INK2})]],
                      size=12, name="Team"))

# ================================================================== 2 BACKGROUND
s, m = new_slide(2, T("BACKGROUND", "LATAR BELAKANG"), T("Food waste is Indonesia's", "Sisa makanan:"),
                 T("largest organic waste stream.", "sampah organik terbesar Indonesia."))
STATS = [("53 kg", T("household food waste per capita per year", "sisa makanan rumah tangga per kapita per tahun"), "UNEP (2024)", GREEN),
         ("14.7 Mt", T("household food waste per year, Indonesia", "sisa makanan rumah tangga per tahun, Indonesia"), "UNEP (2024)", TEAL),
         ("≈40%", T("of municipal solid waste is food waste", "sampah kota berupa sisa makanan"), "SIPSN KLHK (2023)", AMBER),
         ("1.7 Gt", T("CO₂e from food loss & waste, 2000–2019", "CO₂e dari susut & sisa pangan, 2000–2019"), "Bappenas (2021)", RED)]
for i, (big, lab, src, col) in enumerate(STATS):
    x = M + i * 3.07
    m["anim"].append(card(s, x, 1.95, 2.8, 1.75, name=f"Stat_{i}_card"))
    m["anim"].append(text(s, x + 0.22, 2.08, 2.4, 1.6, [[(dec(big), {"size": 30, "bold": True, "color": col})], [(lab, {"size": 11.5})],
                                                       [(src, {"size": 9, "color": MUTED})]], spacing=1.0, name=f"Stat_{i}"))
m["anim"].append(card(s, M, 4.0, 5.9, 2.55, "aura", name="Feed_card"))
m["anim"].append(text(s, M + 0.25, 4.15, 5.5, 0.4, T("Why anaerobic digestion?", "Mengapa pencernaan anaerob?"), size=15, bold=True, color=DEEP, name="Feed_h"))
ntable_ = ntable(s, M + 0.25, 4.6, 5.4, [[T("Property of food waste", "Sifat sisa makanan"), T("Typical value", "Nilai tipikal"), T("Source", "Sumber")],
                                        ["Total solids (TS)", "23–27 %", "Xu et al. (2018)"],
                                        ["Volatile solids (VS/TS)", "90–95 %", "Xu et al. (2018)"],
                                        [T("Biochemical CH₄ potential", "Potensi CH₄ biokimia"), "330–490 mL CH₄/g VS", "Zhang et al. (2007)"],
                                        [T("Theoretical (Buswell)", "Teoretis (Buswell)"), "≈ 643 mL CH₄/g VS", "Buswell & Mueller (1952)"]],
                 [2.1, 1.65, 1.65], size=10, name="Feed_tbl")
m["anim"].append(ntable_)
m["anim"].append(card(s, 6.85, 4.0, 5.85, 2.55, name="RQ_card"))
m["anim"].append(text(s, 7.1, 4.15, 5.4, 0.4, T("Research questions", "Pertanyaan penelitian"), size=15, bold=True, color=DEEP, name="RQ_h"))
bullets(s, m, 7.1, 4.6, 5.4, 1.9, [
    ("RQ1 ", T("What digester volume, OLR and HRT keep household food-waste digestion stable?", "Berapa volume, OLR, dan HRT agar pencernaan sisa makanan rumah tangga stabil?")),
    ("RQ2 ", T("How much biogas, cooking time and avoided CO₂e can one unit deliver?", "Berapa biogas, waktu masak, dan CO₂e terhindar dari satu unit?")),
    ("RQ3 ", T("Can sensor feedback make the process usable for non-experts (SUS ≥ 68)?", "Apakah umpan balik sensor membuat proses mudah dipakai awam (SUS ≥ 68)?"))], size=12, name="RQ")
cite(s, m, "UNEP Food Waste Index Report 2024; SIPSN KLHK 2023; Bappenas 2021 FLW study; Zhang et al. 2007, Bioresour. Technol. 98:929; Xu et al. 2018, Bioresour. Technol. 247:1047.")

# ================================================================== 3 STATE OF THE ART
s, m = new_slide(3, T("STATE OF THE ART", "STATE OF THE ART"), T("Current options trade energy,", "Solusi kini korbankan"),
                 T("space or visibility.", "energi, ruang, atau pantauan."))
SOTA = [[T("Approach", "Pendekatan"), T("Footprint / scale", "Ukuran / skala"), T("Input", "Input"), T("Energy balance", "Neraca energi"),
         T("Monitoring", "Pemantauan"), T("Main limitation", "Keterbatasan utama")],
        [T("Electric food recycler (Lomi)", "Food recycler elektrik (Lomi)"), T("countertop", "di meja dapur"), T("food scraps", "sisa makanan"),
         T("consumes 0.6–1.0 kWh/cycle", "memakai 0,6–1,0 kWh/siklus"), "app", T("no energy recovery", "tanpa pemulihan energi")],
        ["HomeBiogas 2", "2.1 × 1.15 m, 2.0 m³", T("4 L FW + 4 L water/d", "4 L SM + 4 L air/hari"), T("up to 2 h cooking/d", "hingga 2 jam masak/hari"),
         T("visual", "visual"), T("outdoor, 2.4 m² footprint", "luar ruang, tapak 2,4 m²")],
        [T("Fixed-dome digester (BIRU)", "Digester kubah tetap (BIRU)"), T("masonry, outdoor", "permanen, luar ruang"), T("livestock manure", "kotoran ternak"),
         T("household cooking gas", "gas masak rumah tangga"), T("none", "tidak ada"), T("needs manure; 25,157 units (2021)", "butuh ternak; 25.157 unit (2021)")],
        [T("IoT digester research", "Riset digester IoT"), T("pilot / lab", "pilot / lab"), T("mixed organics", "organik campuran"),
         T("measured, not productised", "terukur, belum jadi produk"), "ESP32: T, pH, CH₄, H₂S", T("not a household device", "bukan perangkat rumah tangga")],
        ["ARUNA", "0.44 × 0.44 m, 100 L", T("1.5 kg FW + 1.5 L water/d", "1,5 kg SM + 1,5 L air/hari"),
         T(f"≈{f0(NM['cook'])} min cooking/d (predicted)", f"≈{f0(NM['cook'])} menit masak/hari (prediksi)"),
         T("NDIR CH₄, pH, T, level + interlocks", "NDIR CH₄, pH, T, level + interlock"), T("to be field-validated", "perlu validasi lapangan")]]
tb = ntable(s, M, 1.95, 12.1, SOTA, [2.45, 1.7, 1.85, 2.2, 2.0, 1.9], size=10, name="SOTA", row_h=0.52)
for j in range(6):
    c = tb.table.cell(5, j)
    c.fill.fore_color.rgb = SOFT
    for r in c.text_frame.paragraphs[0].runs:
        r.font.bold, r.font.color.rgb = True, DEEP
m["anim"].append(tb)
m["anim"].append(card(s, M, 5.25, 12.1, 1.25, "aura", name="Gap_card"))
m["anim"].append(text(s, M + 0.3, 5.38, 11.5, 1.05, [[(T("Research gap  ", "Celah riset  "), {"bold": True, "color": GREEN, "size": 15})],
    [(T("No household option combines food-waste input, energy recovery, indoor footprint and real-time process monitoring. Low-cost MQ-type gas sensors (300–10,000 ppm) cannot quantify biogas CH₄ (50–70 % vol); an NDIR sensor (0–100 % vol) is required.",
        "Belum ada opsi rumah tangga yang menggabungkan input sisa makanan, pemulihan energi, tapak dalam ruang, dan pemantauan proses real-time. Sensor MQ murah (300–10.000 ppm) tidak dapat mengukur CH₄ biogas (50–70 % vol); dibutuhkan sensor NDIR (0–100 % vol)."), {"size": 12, "color": INK2})]],
    spacing=1.05, name="Gap"))
cite(s, m, "Lomi (2024) energy use; HomeBiogas (2024) HB2 specifications; ESDM/Hivos BIRU (2021); Prasetyono et al. 2026, J. Mechatron. Electr. Power Veh. Technol. 17(1):29; Rajendran et al. 2012, Energies 5:2911; Winsen MQ-4 & MH-440D datasheets.")

# ================================================================== 4 SOLUTION + SDGs
s, m = new_slide(4, T("SOLUTION", "SOLUSI"), T("Shred, stir, digest, measure:", "Cacah, aduk, cerna, ukur:"), T("three outputs.", "tiga hasil."))
figure(s, m, os.path.join(SF, "sci_pfd.png"), M, 1.95, w=7.4, name="PFD")
OUTS = [(T("Biogas", "Biogas"), T(f"{f0(NM['bg'])} L/d · {f1(NM['kwh'])} kWh/d", f"{f0(NM['bg'])} L/hari · {dec(f1(NM['kwh']))} kWh/hari"), BLUE),
        (T("Liquid digestate", "Digestat cair"), T("N-rich fertilizer, dilute 1:10", "pupuk kaya N, encerkan 1:10"), GREEN),
        (T("Solid digestate", "Digestat padat"), T("compost drawer", "laci kompos"), AMBER)]
for i, (a, b, col) in enumerate(OUTS):
    y = 1.9 + i * 0.78
    m["anim"].append(card(s, 8.35, y, 4.35, 0.62, name=f"Out_{i}_card"))
    m["anim"].append(dot(s, 8.5, y + 0.11, 0.4, col, str(i + 1), size=11, name=f"Out_{i}_dot"))
    m["anim"].append(text(s, 9.05, y + 0.06, 3.6, 0.55, [[(a, {"bold": True, "size": 12.5})], [(b, {"size": 10.5, "color": INK2})]], name=f"Out_{i}"))
SDG = [("7", "7.1", T(f"clean cooking fuel ≈ {f1(NM['lpg'])} kg LPG-eq/month", f"bahan bakar masak bersih ≈ {dec(f1(NM['lpg']))} kg LPG-eq/bulan")),
       ("11", "11.6", T(f"{FWYR} t/yr diverted from municipal waste", f"{dec(FWYR)} t/thn dialihkan dari sampah kota")),
       ("12", "12.3", T("food waste valorised at source", "sisa makanan diolah di sumber")),
       ("13", "13.2", T(f"{CO2N} t CO₂e/yr avoided (potential)", f"{dec(CO2N)} t CO₂e/thn terhindar (potensi)"))]
for i, (n, tg, d) in enumerate(SDG):
    x, y = 8.35 + (i % 2) * 2.2, 4.32 + (i // 2) * 1.1
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.1), Inches(0.98))
    b.adjustments[0] = 0.12
    b.fill.solid()
    b.fill.fore_color.rgb = SDGC[n]
    b.line.fill.background()
    b.name = f"SDG_{i}"
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.05)
    tf.word_wrap = True
    for k, (t_, sz, bd) in enumerate(((f"SDG {tg}", 13, True), (d, 9, False))):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t_
        r.font.size, r.font.bold, r.font.name = Pt(sz), bd, FONT
        r.font.color.rgb = WHITE
    m["anim"].append(b)
m["anim"].append(card(s, M, 5.55, 7.4, 0.85, "aura", name="Hyp_card"))
m["anim"].append(text(s, M + 0.2, 5.65, 7.1, 0.75, [[("H1 ", {"bold": True, "color": AMBER}),
    (T(f"At OLR ≤ {dec(f1(NM['olr']))} kg VS/m³·d and HRT ≈ {f0(NM['hrt'])} d: CH₄ ≥ 50 %, ≥ 0.30 m³ CH₄/kg VS, FOS/TAC < 0.4.   ",
       f"Pada OLR ≤ {dec(f1(NM['olr']))} kg VS/m³·hari dan HRT ≈ {f0(NM['hrt'])} hari: CH₄ ≥ 50 %, ≥ 0,30 m³ CH₄/kg VS, FOS/TAC < 0,4.   "), {"color": INK2}),
    ("H2 ", {"bold": True, "color": AMBER}), (T("SUS ≥ 68.", "SUS ≥ 68."), {"color": INK2})]], size=11, name="Hyp"))
cite(s, m, "Design calculation (this work) using Zhang et al. 2007; Xu et al. 2018; IPCC 2006 Vol. 5 Ch. 3; IPCC AR5 GWP100 = 28; UN A/RES/70/1 (SDG targets).")

# ================================================================== 5 THE PRODUCT
s, m = new_slide(5, T("THE PRODUCT", "PRODUK"), T("A 0.19 m² digester", "Digester 0,19 m²"), T("with a digital twin.", "dengan kembaran digital."))
cv, cw, ch = fit_image(os.path.join(A, "hw_outer.png"), 3.3, 4.55, crop=(0.18, 0.12, 0.82, 0.92), key="hw_outer_sci")
m["anim"].append(card(s, M, 1.95, 3.6, 4.6, "aura", name="HW_card"))
c = pic(s, cv, M + 0.15, 1.98, cw, ch, name="FX_FLOAT_HW")
m["anim"].append(c)
m["float"].append(c)
SPEC = [[T("Hardware", "Hardware"), T("Value", "Nilai")],
        [T("Footprint × height", "Tapak × tinggi"), "44 × 44 × 112 cm"],
        [T("Working volume", "Volume kerja"), "100 L"],
        [T("Nominal / peak load", "Beban nominal / puncak"), T("1.5 / 3.0 kg FW/d", "1,5 / 3,0 kg SM/hari")],
        ["HRT / OLR", T(f"{f0(NM['hrt'])} d / {f1(NM['olr'])} kg VS/m³·d", f"{f0(NM['hrt'])} hari / {dec(f1(NM['olr']))}")],
        [T("Controller", "Kontroler"), "ESP32-S3 (Wi-Fi, BLE)"],
        [T("Process sensors", "Sensor proses"), "T, pH, NDIR CH₄, level, mass"],
        [T("Safety", "Keamanan"), T("MQ-4 leak alarm, valve", "alarm bocor MQ-4, katup")]]
m["anim"].append(ntable(s, 4.45, 1.95, 3.95, SPEC, [1.6, 2.35], size=10, name="SpecHW", row_h=0.4))
figure(s, m, os.path.join(SF, "fig_architecture.png"), 8.75, 2.0, w=3.95, name="Arch")
SW = [[T("Software", "Software"), T("Implementation", "Implementasi")],
      [T("Telemetry", "Telemetri"), T("15-min sampling → Firebase RTDB", "sampling 15 mnt → Firebase RTDB")],
      [T("Interface", "Antarmuka"), T("PWA, live at aruna-id.web.app", "PWA, live di aruna-id.web.app")],
      [T("Guidance", "Panduan"), T("interlock alerts + AI educator", "alert interlock + AI edukator")],
      [T("Setup", "Setup"), T("BLE provisioning, offline", "provisioning BLE, offline")]]
m["anim"].append(ntable(s, 8.75, 4.25, 3.95, SW, [1.25, 2.7], size=10, name="SpecSW", row_h=0.42))
cite(s, m, "This work (3D model & live app); Espressif ESP32-S3 datasheet; Google Firebase Realtime Database documentation.")

# ================================================================== 6 THE HARDWARE
s, m = new_slide(6, T("THE HARDWARE", "HARDWARE"), T("Gravity-fed layout,", "Alir gravitasi,"), T("instrumented for process control.", "berinstrumen kontrol proses."))
xr = os.path.join(A, "hw_xray.png")
PTS = json.load(open(os.path.join(A, "hw_points.json")))
CROP = (0.2, 0.1, 0.8, 0.95)
cv, cw, ch = fit_image(xr, 3.6, 4.75, crop=CROP, key="hw_xray_sci")
IX, IY = 2.95, 1.85
m["anim"].append(pic(s, cv, IX, IY, cw, ch, name="Xray"))
LBL = [("pencacah", T("Twin-shaft shredder", "Pencacah dua poros")), ("loadcell", T("Load cell (HX711)", "Load cell (HX711)")),
       ("layar", T("7″ panel + ESP32-S3", "Panel 7″ + ESP32-S3")), ("pengaduk", T("Stirrer, 5 min / 2 h", "Pengaduk, 5 mnt / 2 jam")),
       ("keran", T("Liquid digestate tap", "Keran digestat cair")), ("kompos", T("Solid digestate drawer", "Laci digestat padat"))]
LBL.sort(key=lambda r: PTS[r[0]][1])
for i, (k, lab) in enumerate(LBL):
    u, v = PTS[k]
    px = IX + (u - CROP[0]) / (CROP[2] - CROP[0]) * cw
    py = IY + (v - CROP[1]) / (CROP[3] - CROP[1]) * ch
    ty = 2.0 + i * 0.75
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.75), Inches(ty + 0.17), Inches(px), Inches(py))
    ln.line.color.rgb = GREEN
    ln.line.width = Pt(1.25)
    ln.name = f"Line_{k}"
    d = dot(s, px - 0.06, py - 0.06, 0.12, GREEN, name=f"FX_PULSE_Pt_{k}")
    m["anim"] += [text(s, M, ty, 2.12, 0.4, lab, size=11, bold=True, align=PP_ALIGN.RIGHT, name=f"Lab_{k}"), ln, d]
INST = [[T("Parameter", "Parameter"), T("Sensor", "Sensor"), T("Range", "Rentang"), T("Accuracy / role", "Akurasi / fungsi")],
        [T("Temperature", "Suhu"), "DS18B20", "−55–125 °C", "±0.5 °C"],
        [T("CH₄ in biogas", "CH₄ biogas"), "NDIR MH-440D", "0–100 % vol", "T90 < 30 s"],
        [T("CH₄ leak", "Bocor CH₄"), "MQ-4", "300–10,000 ppm", T("alarm 10 % LEL", "alarm 10 % LEL")],
        ["pH", T("glass probe", "probe kaca"), "0–14", T("3-pt calibration", "kalibrasi 3 titik")],
        [T("Level", "Level"), T("ultrasonic", "ultrasonik"), T("0–100 % fill", "0–100 % isi"), "≈ ±1 cm"],
        [T("Input mass", "Massa input"), "load cell + HX711", "0–5 kg", "24-bit ADC"],
        [T("Gas volume", "Volume gas"), T("diaphragm meter", "meter diafragma"), "L/d", T("reference", "referensi")]]
m["anim"].append(ntable(s, 6.85, 1.95, 5.85, INST, [1.35, 1.55, 1.45, 1.5], size=10, name="Inst", row_h=0.42))
m["anim"].append(card(s, 6.85, 5.5, 5.85, 1.0, "aura", name="Int_card"))
m["anim"].append(text(s, 7.05, 5.58, 5.5, 0.9, [[(T("Safety & process interlocks  ", "Interlock keamanan & proses  "), {"bold": True, "color": GREEN})],
    [(T("leak > 5,000 ppm → close valve · pH < 6.5 → pause feeding · T < 30 °C → insulation alert · level > 90 % → harvest prompt",
        "bocor > 5.000 ppm → katup tutup · pH < 6,5 → hentikan umpan · T < 30 °C → alert insulasi · level > 90 % → panen"), {"size": 11, "color": INK2})]],
    size=12, name="Int"))
cite(s, m, "Maxim/Analog Devices DS18B20 datasheet; Winsen MH-440D (NDIR, 0–100 % vol) & MQ-4 (300–10,000 ppm) datasheets; CH₄ LEL 5 % vol.")

# ================================================================== 7 THE SOFTWARE
s, m = new_slide(7, T("THE SOFTWARE", "APLIKASI"), T("From raw signals", "Dari sinyal mentah"), T("to plain-language decisions.", "ke keputusan yang jelas."))
figure(s, m, os.path.join(SF, "sci_control.png"), M, 1.95, h=4.55, name="Control")
for i, key in enumerate(("home", "monitor", "chat")):
    im, iw, ih = fit_image(os.path.join(DESIGN, "app", "screens", f"{key}.png"), 1.95, 4.2, key=f"scr_sci_{key}")
    x = 5.55 + i * 2.2
    m["anim"].append(card(s, x - 0.05, 1.95, iw + 0.1, ih + 0.1, r=0.25, name=f"Ph_{i}_card"))
    p = pic(s, im, x, 2.0, iw, ih, name=f"FX_FLOAT_Ph_{i}" if i == 1 else f"Ph_{i}")
    m["anim"].append(p)
    if i == 1:
        m["float"].append(p)
    m["anim"].append(text(s, x, 6.27, iw, 0.3, [T("Status", "Status"), T("Trends + ideal band", "Tren + zona ideal"), T("AI educator", "AI edukator")][i],
                          size=10, bold=True, color=INK2, align=PP_ALIGN.CENTER, name=f"PhL_{i}"))
cite(s, m, "This work (rule-based controller, live PWA). Rule thresholds: pH 6.5–7.5 and FOS/TAC < 0.4 (Chen et al. 2008, Bioresour. Technol. 99:4044; Liu et al. 2021, Appl. Sci. 11:11843).")

# ================================================================== 8 USE CASE
s, m = new_slide(8, T("USE CASE", "USE CASE"), T("One day of operation,", "Satu hari operasi,"), T("quantified.", "terkuantifikasi."))
STEPS = [(96, T("Feed", "Isi"), T("1.5 kg FW + 1.5 L water; weighed & shredded", "1,5 kg SM + 1,5 L air; ditimbang & dicacah")),
         (165, T("Digest", "Cerna"), T(f"OLR {f1(NM['olr'])} kg VS/m³·d, stir 5 min / 2 h", f"OLR {dec(f1(NM['olr']))}, aduk 5 mnt / 2 jam")),
         (250, T("Monitor", "Pantau"), T("15-min telemetry; alert if pH < 6.5", "telemetri 15 mnt; alert jika pH < 6,5")),
         (330, T("Fertilize", "Pupuk"), T("liquid digestate diluted 1:10", "digestat cair diencerkan 1:10")),
         (430, T("Cook", "Masak"), T(f"{f0(NM['bg'])} L biogas ≈ {f0(NM['cook'])} min at 0.30 m³/h", f"{f0(NM['bg'])} L biogas ≈ {f0(NM['cook'])} mnt pada 0,30 m³/jam"))]
for i, (fr, a, b) in enumerate(STEPS):
    x = M + i * 2.46
    im, iw, ih = fit_image(frame(fr), 2.25, 3.4, crop=(0, 0.1, 1, 0.68), key=f"uc_sci_{fr}")
    m["anim"].append(card(s, x - 0.04, 1.95, 2.3, 4.6, r=0.22, name=f"UC_{i}_card"))
    m["anim"].append(pic(s, im, x + 0.025, 2.02, 2.16, None, name=f"UC_{i}"))
    m["anim"].append(dot(s, x + 0.12, 2.12, 0.42, GREEN, str(i + 1), size=12, name=f"UC_{i}_dot"))
    m["anim"].append(text(s, x + 0.1, 5.3, 2.1, 1.2, [[(a, {"bold": True, "size": 14})], [(b, {"size": 10.5, "color": INK2})]], name=f"UC_{i}_t"))
cite(s, m, "Design calculation (this work); stove demand 0.2–0.45 m³/h (SSWM/Akvopedia, Biogas combustion, 2019).")

# ================================================================== 9 METHOD
s, m = new_slide(9, T("RESEARCH METHOD", "METODE PENELITIAN"), T("R&D (4D model) with", "R&D (model 4D) dengan"), T("quantitative exit criteria.", "kriteria lulus kuantitatif."))
figure(s, m, os.path.join(SF, "fig_method_4d.png"), M, 1.95, w=6.3, name="M4D")
EQ = [(T("Mass balance", "Neraca massa"), "OLR = F·TS·(VS/TS) / V_w_      HRT = V_w_ / Q"),
      (T("Kinetics (Gompertz)", "Kinetika (Gompertz)"), "M(t) = P·exp{−exp[(R_m_·e/P)(λ − t) + 1]}"),
      (T("Theoretical CH₄ (Buswell)", "CH₄ teoretis (Buswell)"), "C_c_H_h_O_o_N_n_ + H₂O → CH₄ + CO₂ + NH₃"),
      (T("Energy & cooking", "Energi & memasak"), "E = V_CH4_ · 35.8 MJ/m³ ;  t_cook_ = V_biogas_ / q_stove_"),
      (T("Avoided landfill CH₄ (IPCC)", "CH₄ TPA terhindar (IPCC)"), "W · DOC · DOC_f_ · MCF · F · 16/12 · (1 − OX)")]


def eq_runs(e, size=14):
    import re as _re
    out = []
    for part in _re.split(r"(_[^_]+_)", e):
        if not part:
            continue
        if part.startswith("_") and part.endswith("_"):
            out.append((part[1:-1], {"font": "Cambria Math", "size": size, "sub": True}))
        else:
            out.append((part, {"font": "Cambria Math", "size": size}))
    return out
m["anim"].append(card(s, 7.15, 1.95, 5.55, 2.85, "aura", name="Eq_card"))
m["anim"].append(text(s, 7.35, 2.08, 5.2, 2.7, [line for a, b in EQ for line in ([(a, {"bold": True, "color": TEAL, "size": 10})], eq_runs(b, 14.5))],
                      spacing=1.0, name="Eq"))
DES = [[T("Experiment", "Eksperimen"), T("Design", "Desain"), T("Standard / method", "Standar / metode"), T("Output", "Luaran")],
       ["BMP assay", T("ISR ≥ 2, 37 °C, triplicate", "ISR ≥ 2, 37 °C, triplo"), "VDI 4630", "P, Rm, λ"],
       [T("Semi-continuous", "Semi-kontinu"), T("OLR 2.0 / 3.4 / 4.5 × 3 reps, 2 HRT", "OLR 2,0 / 3,4 / 4,5 × 3 ulangan, 2 HRT"), "APHA 2540 G; FOS/TAC", T("yield, stability", "yield, stabilitas")],
       [T("Field test", "Uji lapangan"), T("[n] households × 30 d", "[n] rumah tangga × 30 hari"), T("telemetry + logbook", "telemetri + logbook"), T("CH₄ %, L/d, kg/d", "CH₄ %, L/hari, kg/hari")],
       [T("Usability", "Usability"), T("post-test questionnaire", "kuesioner pasca uji"), "SUS (Brooke 1996)", T("score 0–100", "skor 0–100")]]
m["anim"].append(ntable(s, M, 4.98, 7.7, DES, [1.55, 2.6, 1.95, 1.6], size=10, name="Design", row_h=0.33))
m["anim"].append(card(s, 8.55, 4.98, 4.15, 1.58, name="Stat_card"))
bullets(s, m, 8.72, 5.06, 3.9, 1.5, [(T("Kinetics: ", "Kinetika: "), T("non-linear least squares, R², RMSE", "least squares non-linear, R², RMSE")),
                                     (T("Comparison: ", "Perbandingan: "), T("one-way ANOVA + Tukey, α = 0.05", "ANOVA satu arah + Tukey, α = 0,05")),
                                     (T("Biodegradability: ", "Biodegradabilitas: "), "BMP(exp) / BMP(theoretical)")], size=11, name="Stats")
cite(s, m, "Thiagarajan et al. 1974; Buswell & Mueller 1952, Ind. Eng. Chem. 44:550; Zwietering et al. 1990, Appl. Environ. Microbiol. 56:1875; VDI 4630 (2016); APHA 2540 (2017); Brooke 1996.")

# ================================================================== 10 FIELD TESTING
s, m = new_slide(10, T("FIELD TESTING", "UJI LAPANGAN"), T("Measurement plan with", "Rencana pengukuran dengan"), T("acceptance criteria.", "kriteria penerimaan."))
illustrative(s, m, M, 1.62)
PLAN = [[T("Parameter", "Parameter"), T("Instrument / method", "Instrumen / metode"), T("Frequency", "Frekuensi"), T("Acceptance", "Kriteria"), T("Basis", "Dasar")],
        [T("Temperature", "Suhu"), "DS18B20", T("15 min", "15 menit"), "30–38 °C", "Weiland (2010)"],
        ["pH", T("calibrated probe", "probe terkalibrasi"), T("15 min", "15 menit"), "6.8–7.5", "Chen et al. (2008)"],
        ["FOS/TAC", T("two-point titration", "titrasi dua titik"), T("weekly", "mingguan"), "< 0.4", "Liu et al. (2021)"],
        [T("CH₄ content", "Kadar CH₄"), "NDIR + [ref. analyser]", T("continuous", "kontinu"), "≥ 50 % vol", "Xu et al. (2018)"],
        [T("Biogas volume", "Volume biogas"), T("diaphragm gas meter", "meter gas diafragma"), T("daily", "harian"), f"≥ {f0(NM['bg']*0.8)} L/d", T("this work (−20 %)", "kajian ini (−20 %)")],
        ["TS / VS", "APHA 2540 G", T("weekly", "mingguan"), T("VS removal ≥ 70 %", "VS tersisih ≥ 70 %"), "Mata-Alvarez (2000)"],
        [T("Digestate N, P, K", "N, P, K digestat"), T("lab analysis", "analisis lab"), T("end of cycle", "akhir siklus"), "[standard]", T("national fertilizer standard", "standar pupuk nasional")],
        [T("Usability", "Usability"), "SUS", T("end of test", "akhir uji"), "≥ 68", "Bangor et al. (2009)"]]
m["anim"].append(ntable(s, M, 2.12, 8.2, PLAN, [1.55, 2.0, 1.25, 1.6, 1.8], size=10, name="Plan", row_h=0.47))
im, iw, ih = fit_image(frame(250), 3.6, 3.9, crop=(0.0, 0.18, 1.0, 0.75), key="ft_sci")
m["anim"].append(card(s, 9.05, 2.12, 3.65, 4.25, "aura", r=0.25, name="FT_card"))
m["anim"].append(pic(s, im, 9.15, 2.22, 3.45, None, name="FT_img"))
m["anim"].append(text(s, 9.1, 6.38, 3.6, 0.3, T("Rendered scenario; replace with field photos.", "Skenario render; ganti dengan foto lapangan."),
                      size=8.5, color=MUTED, align=PP_ALIGN.CENTER, name="FT_note"))
cite(s, m, "Weiland 2010; Chen et al. 2008; Liu et al. 2021, Appl. Sci. 11(24):11843; Xu et al. 2018; Mata-Alvarez et al. 2000, Bioresour. Technol. 74:3; APHA 2540 (2017); Bangor et al. 2009, J. Usability Stud. 4:114.")

# ================================================================== 11 RESULTS & DISCUSSION
s, m = new_slide(11, T("RESULT & DISCUSSION", "HASIL & PEMBAHASAN"), T("Design predictions", "Prediksi desain"), T("against the literature.", "dibandingkan literatur."))
m["anim"].append(pill(s, M, 1.62, T("PREDICTED FROM DESIGN CALCULATION · measured data pending", "PREDIKSI DARI PERHITUNGAN DESAIN · data ukur menyusul"),
                      fill=RGBColor(0xFB, 0xF0, 0xDF), color=AMBER, size=9, name="Tag_Pred"))
RES = [[T("Parameter", "Parameter"), T("Household of 4", "RT 4 orang"), T("Nominal", "Nominal"), T("Peak", "Puncak")],
       [T("Input (kg FW/d)", "Input (kg SM/hari)"), f"{H4['fw']:.2f}", "1.5", "3.0"],
       ["OLR (kg VS/m³·d)", f1(H4["olr"]), f1(NM["olr"]), f1(MX["olr"])],
       [T("HRT (d)", "HRT (hari)"), f0(H4["hrt"]), f0(NM["hrt"]), f0(MX["hrt"])],
       [T("Biogas (L/d)", "Biogas (L/hari)"), f0(H4["bg"]), f0(NM["bg"]), f0(MX["bg"])],
       [T("Cooking (min/d)", "Masak (mnt/hari)"), f0(H4["cook"]), f0(NM["cook"]), f0(MX["cook"])],
       [T("LPG-eq (kg/month)", "Setara LPG (kg/bln)"), f1(H4["lpg"]), f1(NM["lpg"]), f1(MX["lpg"])],
       [T("Avoided (t CO₂e/yr)", "Terhindar (t CO₂e/thn)"), f"{H4['co2e']:.2f}", f"{NM['co2e']:.2f}", f"{MX['co2e']:.2f}"]]
if LANG == "id":
    RES = [[dec(c) if i else c for i, c in enumerate(r)] for r in RES]
m["anim"].append(ntable(s, M, 2.12, 5.4, RES, [2.1, 1.15, 1.05, 1.1], size=10.5, hl_col=2, name="Res", row_h=0.4))
figure(s, m, os.path.join(SF, "sci_gompertz.png"), 6.35, 2.12, w=6.3, name="Gomp")
bullets(s, m, M, 5.45, 12.1, 1.25, [
    (T("Biodegradability ≈ 61 %: ", "Biodegradabilitas ≈ 61 %: "), T("BMP 390 vs Buswell 643 mL CH₄/g VS, within reported food-waste range (Zhang et al., 2007).", "BMP 390 vs Buswell 643 mL CH₄/g VS, dalam rentang literatur (Zhang et al., 2007).")),
    (T("Stability: ", "Stabilitas: "), T(f"nominal OLR {f1(NM['olr'])} is within stable ranges; peak OLR {f1(MX['olr'])} at HRT {f0(MX['hrt'])} d risks VFA accumulation → pH interlock (Chen et al., 2008).",
                                       f"OLR nominal {dec(f1(NM['olr']))} dalam rentang stabil; OLR puncak {dec(f1(MX['olr']))} pada HRT {f0(MX['hrt'])} hari berisiko akumulasi VFA → interlock pH (Chen et al., 2008).")),
    (T("Scale: ", "Skala: "), T(f"a household of 4 yields ≈ {f0(H4['cook'])} min/d; nominal load fits extended families or small eateries.", f"rumah tangga 4 orang ≈ {f0(H4['cook'])} mnt/hari; beban nominal cocok untuk keluarga besar atau warung kecil."))],
    size=11.5, name="Disc")
cite(s, m, "Design calculation (this work); Gompertz parameters from literature ranges, to be fitted (Zwietering et al. 1990); Zhang et al. 2007; Chen et al. 2008; UNEP 2024 (53 kg/cap/yr).")

# ================================================================== 12 FUTURE DEVELOPMENT
s, m = new_slide(12, T("FUTURE DEVELOPMENT", "PENGEMBANGAN"), T("From prototype", "Dari prototipe"), T("to verified impact.", "ke dampak terverifikasi."))
figure(s, m, os.path.join(SF, "sci_sensitivity.png"), M, 1.95, w=5.6, name="Sens")
ROAD = [(T("0–6 months", "0–6 bulan"), T("Validate", "Validasi"), [T("BMP + semi-continuous tests; fit P, Rm, λ", "Uji BMP + semi-kontinu; fit P, Rm, λ"),
                                                                     T("calibrate NDIR vs gas analyser", "kalibrasi NDIR vs gas analyser")], GREEN),
        (T("6–12 months", "6–12 bulan"), T("Optimise", "Optimasi"), [T("co-digestion to C/N 20–30", "ko-digesti ke C/N 20–30"),
                                                                       T("H₂S scrubber; insulation for T < 25 °C", "scrubber H₂S; insulasi untuk T < 25 °C")], TEAL),
        (T("12–36 months", "12–36 bulan"), T("Scale", "Skala"), [T("on-device AI waste recognition", "AI pengenal sampah di perangkat"),
                                                                   T("MBG school-kitchen pilots with MRV", "pilot dapur MBG dengan MRV")], DEEP)]
for i, (when, ttl, items, col) in enumerate(ROAD):
    y = 1.95 + i * 1.5
    m["anim"].append(card(s, 6.65, y, 6.05, 1.32, "aura" if i == 0 else "raise", name=f"Road_{i}_card"))
    m["anim"].append(dot(s, 6.82, y + 0.18, 0.46, col, str(i + 1), size=13, name=f"FX_PULSE_Road_{i}" if i == 0 else f"Road_{i}_dot"))
    m["anim"].append(text(s, 7.45, y + 0.1, 5.1, 1.2, [[(when + "  ·  ", {"bold": True, "size": 10, "color": col}), (ttl, {"bold": True, "size": 15})]] +
                          [[("● ", {"color": col, "size": 10}), (it, {"size": 11.5, "color": INK2})] for it in items], spacing=1.05, name=f"Road_{i}"))
cite(s, m, "Sensitivity analysis (this work); C/N 20–30 and H₂S removal (Weiland 2010, Appl. Microbiol. Biotechnol. 85:849; Mata-Alvarez et al. 2000).")

# ================================================================== 13 CONCLUSION
s, m = new_slide(13, T("CONCLUSION", "KESIMPULAN"), T("Answers to", "Jawaban atas"), T("the research questions.", "pertanyaan penelitian."))
CON = [("RQ1", T(f"A 100 L digester at OLR {f1(NM['olr'])} kg VS/m³·d and HRT ≈ {f0(NM['hrt'])} d is predicted to be stable for 1.5 kg/d; peaks above ≈ 3 kg/d require pausing.",
                 f"Digester 100 L pada OLR {dec(f1(NM['olr']))} kg VS/m³·hari dan HRT ≈ {f0(NM['hrt'])} hari diprediksi stabil untuk 1,5 kg/hari; puncak ≈ 3 kg/hari perlu dijeda.")),
       ("RQ2", T(f"≈ {f0(NM['bg'])} L biogas/d → ≈ {f0(NM['cook'])} min cooking, ≈ {f1(NM['lpg'])} kg LPG-eq/month and ≈ {NM['co2e']:.2f} t CO₂e/yr avoided.",
                 f"≈ {f0(NM['bg'])} L biogas/hari → ≈ {f0(NM['cook'])} menit masak, ≈ {dec(f1(NM['lpg']))} kg LPG-eq/bulan, dan ≈ {dec(CO2N)} t CO₂e/thn terhindar.")),
       ("RQ3", T("NDIR sensing, rule-based interlocks and app guidance close the monitoring gap; SUS to be measured.",
                 "Sensor NDIR, interlock berbasis aturan, dan panduan aplikasi menutup celah pemantauan; SUS akan diukur."))]
for i, (rq, txt) in enumerate(CON):
    y = 1.95 + i * 1.12
    m["anim"].append(card(s, M, y, 7.0, 0.95, name=f"Con_{i}_card"))
    m["anim"].append(dot(s, M + 0.18, y + 0.22, 0.52, GREEN, rq, size=10, name=f"Con_{i}_dot"))
    m["anim"].append(text(s, M + 0.85, y + 0.12, 6.0, 0.8, txt, size=12, color=INK2, name=f"Con_{i}"))
figure(s, m, os.path.join(SF, "sci_ghg.png"), 8.0, 1.95, w=4.7, name="GHG")
import qrcode  # noqa: E402
qr = qrcode.QRCode(border=1, box_size=10)
qr.add_data("https://aruna-id.web.app")
qpath = os.path.join(CACHE, "qr.png")
qr.make_image(fill_color=(16, 32, 26), back_color=(255, 255, 255)).save(qpath)
m["anim"].append(card(s, M, 5.4, 7.0, 1.1, "aura", name="QR_card"))
m["anim"].append(pic(s, qpath, M + 0.15, 5.48, 0.92, 0.92, name="QR"))
m["anim"].append(text(s, M + 1.25, 5.52, 5.6, 0.95, [[(T("Live prototype app  ", "Aplikasi prototipe live  "), {"bold": True, "size": 14}), ("aruna-id.web.app", {"bold": True, "color": GREEN, "size": 13})],
                                                  [(T("Values are design predictions; measured results will replace them after testing.", "Nilai adalah prediksi desain; akan diganti hasil ukur setelah pengujian."), {"size": 10.5, "color": MUTED})]],
                      name="QR_t"))
cite(s, m, "This work; IPCC 2006 Vol. 5 Ch. 3 (FOD defaults); IPCC AR5 (GWP100 = 28).")

# ================================================================== 14 REFERENCES (lampiran)
s, m = new_slide(14, T("REFERENCES", "REFERENSI"), T("Literature", "Literatur"), T("and data sources.", "dan sumber data."))
REFS = [
    "APHA (2017). Standard Methods for the Examination of Water and Wastewater, 23rd ed., Method 2540.",
    "Appels, L. et al. (2008). Prog. Energy Combust. Sci. 34(6), 755–781. doi:10.1016/j.pecs.2008.06.002",
    "Bangor, A., Kortum, P., Miller, J. (2009). J. Usability Stud. 4(3), 114–123.",
    "Bappenas (2021). Food Loss and Waste in Indonesia. Ministry of National Development Planning.",
    "Brooke, J. (1996). SUS: a 'quick and dirty' usability scale. Usability Evaluation in Industry, 189–194.",
    "Buswell, A. M., Mueller, H. F. (1952). Ind. Eng. Chem. 44(3), 550–552. doi:10.1021/ie50507a033",
    "Chen, Y., Cheng, J. J., Creamer, K. S. (2008). Bioresour. Technol. 99(10), 4044–4064. doi:10.1016/j.biortech.2007.01.057",
    "ESDM (2021). Strategi pengembangan biogas menuju target 1 juta biodigester. Ditjen EBTKE.",
    "HomeBiogas (2024). HomeBiogas 2 product specifications. homebiogas.com",
    "IPCC (2006). Guidelines for National GHG Inventories, Vol. 5 Waste, Ch. 3.",
    "IPCC (2014). Climate Change 2014: Synthesis Report (AR5).",
    "KLHK (2023). SIPSN: Komposisi sampah. sipsn.menlhk.go.id",
    "Liu, X. et al. (2021). Appl. Sci. 11(24), 11843. doi:10.3390/app112411843",
    "Lomi (2024). How much energy does Lomi use? lomi.com",
    "Mata-Alvarez, J., Macé, S., Llabrés, P. (2000). Bioresour. Technol. 74(1), 3–16. doi:10.1016/S0960-8524(00)00023-7",
    "Prasetyono, S. et al. (2026). J. Mechatron. Electr. Power Veh. Technol. 17(1), 29–38. doi:10.55981/j.mev.2026.1352",
    "Rajendran, K. et al. (2012). Energies 5(8), 2911–2942. doi:10.3390/en5082911",
    "SSWM / Akvopedia (2019). Biogas combustion. sswm.info",
    "Thiagarajan, S., Semmel, D. S., Semmel, M. I. (1974). Instructional Development for Training Teachers of Exceptional Children. Indiana Univ.",
    "UNEP (2024). Food Waste Index Report 2024. Nairobi.",
    "VDI 4630 (2016). Fermentation of organic materials. Verein Deutscher Ingenieure.",
    "Weiland, P. (2010). Appl. Microbiol. Biotechnol. 85(4), 849–860. doi:10.1007/s00253-009-2246-7",
    "Winsen (n.d.). MH-440D NDIR CH₄ sensor; MQ-4 gas sensor datasheets.",
    "Xu, F. et al. (2018). Bioresour. Technol. 247, 1047–1058. doi:10.1016/j.biortech.2017.09.020",
    "Zhang, R. et al. (2007). Bioresour. Technol. 98(4), 929–935. doi:10.1016/j.biortech.2006.02.039",
    "Zwietering, M. H. et al. (1990). Appl. Environ. Microbiol. 56(6), 1875–1881. doi:10.1128/aem.56.6.1875-1881.1990",
]
half = (len(REFS) + 1) // 2
for k, chunk in enumerate((REFS[:half], REFS[half:])):
    x = M + k * 6.15
    m["anim"].append(card(s, x, 1.85, 5.95, 4.75, name=f"Ref_{k}_card"))
    m["anim"].append(text(s, x + 0.2, 1.97, 5.6, 4.55, [[(f"{i + 1 + k * half}. ", {"bold": True, "color": GREEN}), (r, {})] for i, r in enumerate(chunk)],
                          size=8.6, color=INK2, spacing=1.03, name=f"Ref_{k}"))


# ================================================================== XML: transisi + timing
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = 'xmlns:p="%s" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"' % P

TRANS = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" %s>'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
         'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
         '<p:transition spd="slow" p14:dur="1200"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>') % NS


class Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return self.n


def entr(ids, spid, delay):
    return (f'<p:par><p:cTn id="{ids()}" presetID="42" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{ids()}" dur="650"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{ids()}" dur="650" decel="100000" fill="hold"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="#ppt_y+0.04"/></p:val></p:tav>'
            f'<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'
            f'</p:childTnLst></p:cTn></p:par>')


def floatfx(ids, spid, delay):
    return (f'<p:par><p:cTn id="{ids()}" presetID="0" presetClass="path" presetSubtype="0" fill="hold" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
            f'<p:animMotion origin="layout" path="M 0 0 L 0 -0.014 E" pathEditMode="relative" rAng="0" ptsTypes="AA">'
            f'<p:cBhvr><p:cTn id="{ids()}" dur="2600" accel="50000" decel="50000" autoRev="1" fill="hold" repeatCount="indefinite"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst>'
            f'</p:cBhvr><p:rCtr x="0" y="0"/></p:animMotion></p:childTnLst></p:cTn></p:par>')


def pulse(ids, spid, delay):
    return (f'<p:par><p:cTn id="{ids()}" presetID="6" presetClass="emph" presetSubtype="0" fill="hold" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
            f'<p:animScale><p:cBhvr><p:cTn id="{ids()}" dur="1100" autoRev="1" fill="hold" repeatCount="indefinite" accel="50000" decel="50000"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr><p:by x="112000" y="112000"/></p:animScale>'
            f'</p:childTnLst></p:cTn></p:par>')


def media_play(ids, spid):
    return (f'<p:par><p:cTn id="{ids()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{ids()}" dur="10000" fill="hold"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')


for meta in SLIDES:
    sl = meta["slide"]
    root = sl._element
    for old in root.findall(qn("p:timing")) + root.findall(qn("p:transition")):
        root.remove(old)
    vid = meta["video"].shape_id
    ids = Ids()
    seen, kids, blds = set(), [media_play(ids, vid)], []
    delay = 150
    for sh in meta["anim"]:
        if sh.shape_id in seen:
            continue
        seen.add(sh.shape_id)
        kids.append(entr(ids, sh.shape_id, delay))
        if sh._element.tag == qn("p:sp") and sh.has_text_frame and sh.text_frame.text.strip():
            blds.append(f'<p:bldP spid="{sh.shape_id}" grpId="0" animBg="1"/>')
        delay += 90
    for sh in meta["float"]:
        kids.append(floatfx(ids, sh.shape_id, delay + 200))
    for sh in sl.shapes:
        if sh.name.startswith("FX_PULSE"):
            kids.append(pulse(ids, sh.shape_id, delay + 300))
    timing = (f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
              f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
              f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
              f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
              + "".join(kids) +
              f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
              f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
              f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
              f'<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="{ids()}" repeatCount="indefinite" fill="hold">'
              f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
              f'<p:tgtEl><p:spTgt spid="{vid}"/></p:tgtEl></p:cMediaNode></p:video>'
              f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
              + (f'<p:bldLst>{"".join(blds)}</p:bldLst>' if blds else '') + '</p:timing>')
    clr = root.find(qn("p:clrMapOvr"))
    idx = list(root).index(clr) + 1 if clr is not None else 1
    root.insert(idx, etree.fromstring(TRANS))
    root.insert(idx + 1, etree.fromstring(timing))

out = os.path.join(HERE, f"ARUNA_deck_{LANG}.pptx")
prs.save(out)
print("[DECK]", out, os.path.getsize(out) // 1024, "KB")
